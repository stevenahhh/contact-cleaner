import random
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
import sys
import io

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

def generate_korean_name():
    last_names = ['김', '이', '박', '최', '정', '강', '조', '윤', '장', '임', '한', '오', '서', '신', '권', '황', '안', '송', '류', '홍']
    first_names = ['민수', '지은', '서연', '준호', '하은', '도윤', '서준', '예은', '시우', '지우', '수빈', '은우', '채원', '현우', '지아', '유진', '서현', '민준', '하윤', '연우']
    return random.choice(last_names) + random.choice(first_names)

def generate_phone():
    prefix = random.choice(['010', '011', '016', '017', '018', '019'])
    middle = str(random.randint(1000, 9999))
    last = str(random.randint(1000, 9999))
    return f"{prefix}-{middle}-{last}"

def generate_test_data_hongildong():
    wb = openpyxl.Workbook()
    ws = wb.active
    
    ws.append(['이름', '전화번호'])
    
    base_phones = []
    for i in range(100):
        name = generate_korean_name()
        phone = generate_phone()
        
        if i < 70:
            base_phones.append(phone)
        elif i < 80:
            phone = random.choice(base_phones) if base_phones else phone
        
        ws.append([name, phone])
    
    for row in ws.iter_rows(min_row=1, max_row=1):
        for cell in row:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")
    
    wb.save('테스트_홍길동_주소록.xlsx')
    print(f"[OK] 테스트_홍길동_주소록.xlsx 생성 완료 (100행)")
    return base_phones[:50]

def generate_test_data_kimamugae(shared_phones):
    wb = openpyxl.Workbook()
    ws = wb.active
    
    ws.append(['이름', '연락처'])
    
    for i in range(100):
        name = generate_korean_name()
        
        if i < 30 and shared_phones:
            phone = random.choice(shared_phones)
        elif i < 35:
            phone = generate_phone().replace('-', '')
        elif i < 40:
            phone = generate_phone().replace('-', ' ')
        elif i < 45:
            parts = generate_phone().split('-')
            phone = f"{parts[0]}{parts[1]}{parts[2]}"
        else:
            phone = generate_phone()
        
        ws.append([name, phone])
    
    for row in ws.iter_rows(min_row=1, max_row=1):
        for cell in row:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid")
            cell.alignment = Alignment(horizontal="center")
    
    wb.save('테스트_김아무개_주소록.xlsx')
    print(f"[OK] 테스트_김아무개_주소록.xlsx 생성 완료 (100행)")

print("테스트 데이터 생성 중...")
shared_phones = generate_test_data_hongildong()
generate_test_data_kimamugae(shared_phones)
print("\n생성된 파일:")
print("  1. 테스트_홍길동_주소록.xlsx - 100행 (일부 중복 포함)")
print("  2. 테스트_김아무개_주소록.xlsx - 100행 (다양한 전화번호 형식, 홍길동과 30% 중복)")
print("\n사용 방법:")
print("  - 정리: 각 파일을 개별적으로 정리")
print("  - 대조: 홍길동(비교 주소록) vs 김아무개(대조 대상)")
print("  - 병합: 두 파일을 모두 정리 후 대조 결과를 '주소록_홍길동_대조결과.xlsx', '주소록_김아무개_대조결과.xlsx'로 저장하고 병합")
