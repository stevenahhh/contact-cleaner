import dearpygui.dearpygui as dpg
import threading
from pathlib import Path
import sys
import os
import subprocess

current_dir = Path(__file__).resolve().parent
project_root = current_dir.parent.parent
if str(project_root) not in sys.path:
    sys.path.append(str(project_root))


class ContactCleanerDPG:
    def __init__(self):
        dpg.create_context()
        
        self.clean_file_path = ""
        self.compare_file1_path = ""
        self.compare_file2_path = ""
        self.merge_folder_path = ""
        
        self.processing = False
        
        self._setup_fonts()
        self._setup_theme()
        self._create_ui()
        
    def _setup_fonts(self):
        try:
            from ctypes import windll
            dpi = windll.user32.GetDpiForSystem()
            scale_factor = dpi / 96.0
        except:
            scale_factor = 1.0
        
        font_size = int(15 * scale_factor)
        header_size = int(24 * scale_factor)
        
        with dpg.font_registry():
            malgun_path = "C:/Windows/Fonts/malgun.ttf"
            if os.path.exists(malgun_path):
                with dpg.font(malgun_path, font_size, tag="default_font") as font:
                    dpg.add_font_range_hint(dpg.mvFontRangeHint_Korean)
                    dpg.add_font_range_hint(dpg.mvFontRangeHint_Default)
                
                with dpg.font(malgun_path, header_size, tag="header_font") as font:
                    dpg.add_font_range_hint(dpg.mvFontRangeHint_Korean)
                    dpg.add_font_range_hint(dpg.mvFontRangeHint_Default)
                
                dpg.bind_font("default_font")
    
    def _setup_theme(self):
        with dpg.theme() as global_theme:
            with dpg.theme_component(dpg.mvAll):
                dpg.add_theme_color(dpg.mvThemeCol_WindowBg, (30, 30, 30))
                dpg.add_theme_color(dpg.mvThemeCol_ChildBg, (40, 40, 40))
                dpg.add_theme_color(dpg.mvThemeCol_FrameBg, (50, 50, 50))
                dpg.add_theme_color(dpg.mvThemeCol_FrameBgHovered, (60, 60, 60))
                dpg.add_theme_color(dpg.mvThemeCol_FrameBgActive, (70, 70, 70))
                dpg.add_theme_color(dpg.mvThemeCol_Button, (60, 60, 60))
                dpg.add_theme_color(dpg.mvThemeCol_ButtonHovered, (70, 70, 70))
                dpg.add_theme_color(dpg.mvThemeCol_ButtonActive, (80, 80, 80))
                dpg.add_theme_color(dpg.mvThemeCol_Header, (60, 60, 60))
                dpg.add_theme_color(dpg.mvThemeCol_HeaderHovered, (70, 70, 70))
                dpg.add_theme_color(dpg.mvThemeCol_HeaderActive, (80, 80, 80))
                dpg.add_theme_color(dpg.mvThemeCol_Tab, (50, 50, 50))
                dpg.add_theme_color(dpg.mvThemeCol_TabHovered, (70, 70, 70))
                dpg.add_theme_color(dpg.mvThemeCol_TabActive, (60, 60, 60))
                dpg.add_theme_style(dpg.mvStyleVar_FrameRounding, 5)
                dpg.add_theme_style(dpg.mvStyleVar_WindowRounding, 5)
                dpg.add_theme_style(dpg.mvStyleVar_FramePadding, 8, 6)
        
        dpg.bind_theme(global_theme)
    
    def _create_ui(self):
        with dpg.window(tag="main_window", label="주소록 정리 v1.3.0", width=1150, height=950):
            
            with dpg.tab_bar():
                with dpg.tab(label="주소록 정리"):
                    self._create_clean_tab()
                
                with dpg.tab(label="비교할 주소록"):
                    self._create_compare_tab()
                
                with dpg.tab(label="대조 결과 병합"):
                    self._create_merge_tab()
            
            dpg.add_separator()
            
            with dpg.child_window(height=200, border=True):
                dpg.add_text("로그", tag="log_header")
                dpg.add_separator()
                dpg.add_input_text(
                    tag="log_text",
                    multiline=True,
                    readonly=True,
                    width=-1,
                    height=-1,
                    default_value=""
                )
            
            dpg.add_separator()
            
            with dpg.group(horizontal=True):
                with dpg.child_window(width=200, height=80, border=True):
                    dpg.add_text("상태")
                    dpg.add_text("O: 정상", color=(0, 255, 0))
                    dpg.add_text("△: 주의", color=(255, 255, 0))
                    dpg.add_text("X: 오류", color=(255, 0, 0))
                
                with dpg.child_window(width=-1, height=80, border=True):
                    dpg.add_text("진행 상태", tag="progress_label")
                    dpg.add_progress_bar(tag="progress_bar", default_value=0.0, width=-1)
            
            dpg.add_text("v1.3.0", pos=(1100, 920), color=(136, 136, 136))
    
    def _create_clean_tab(self):
        with dpg.group():
            dpg.add_text("주소록 정리", tag="clean_header")
            dpg.add_separator()
            
            with dpg.child_window(height=250, border=True):
                dpg.add_text("정리할 파일")
                with dpg.group(horizontal=True):
                    dpg.add_button(label="파일 선택", callback=self._select_clean_file, width=150)
                    dpg.add_text("선택된 파일 없음", tag="clean_file_label")
                
                dpg.add_spacer(height=10)
                dpg.add_button(label="정리 시작", callback=self._start_clean, width=150, tag="clean_start_btn")
    
    def _create_compare_tab(self):
        with dpg.group():
            dpg.add_text("주소록 대조", tag="compare_header")
            dpg.add_separator()
            
            with dpg.child_window(height=250, border=True):
                dpg.add_text("대조할 파일 1")
                with dpg.group(horizontal=True):
                    dpg.add_button(label="파일 선택", callback=self._select_compare_file1, width=150)
                    dpg.add_text("선택된 파일 없음", tag="compare_file1_label")
                
                dpg.add_spacer(height=10)
                
                dpg.add_text("대조할 파일 2")
                with dpg.group(horizontal=True):
                    dpg.add_button(label="파일 선택", callback=self._select_compare_file2, width=150)
                    dpg.add_text("선택된 파일 없음", tag="compare_file2_label")
                
                dpg.add_spacer(height=10)
                dpg.add_button(label="대조 시작", callback=self._start_compare, width=150, tag="compare_start_btn")
    
    def _create_merge_tab(self):
        with dpg.group():
            dpg.add_text("대조 결과 병합", tag="merge_header")
            dpg.add_separator()
            
            with dpg.child_window(height=250, border=True):
                dpg.add_text("병합할 폴더")
                with dpg.group(horizontal=True):
                    dpg.add_button(label="폴더 선택", callback=self._select_merge_folder, width=150)
                    dpg.add_text("선택된 폴더 없음", tag="merge_folder_label")
                
                dpg.add_spacer(height=10)
                dpg.add_button(label="병합 시작", callback=self._start_merge, width=150, tag="merge_start_btn")
    
    def _select_clean_file(self):
        dpg.add_file_dialog(
            directory_selector=False,
            show=True,
            callback=self._clean_file_selected,
            file_count=1,
            width=700,
            height=400,
            default_path=str(Path.home() / "Desktop"),
            modal=True
        )
    
    def _clean_file_selected(self, sender, app_data):
        if app_data and app_data['selections']:
            file_path = list(app_data['selections'].values())[0]
            self.clean_file_path = file_path
            dpg.set_value("clean_file_label", Path(file_path).name)
    
    def _select_compare_file1(self):
        dpg.add_file_dialog(
            directory_selector=False,
            show=True,
            callback=self._compare_file1_selected,
            file_count=1,
            width=700,
            height=400,
            default_path=str(Path.home() / "Desktop"),
            modal=True
        )
    
    def _compare_file1_selected(self, sender, app_data):
        if app_data and app_data['selections']:
            file_path = list(app_data['selections'].values())[0]
            self.compare_file1_path = file_path
            dpg.set_value("compare_file1_label", Path(file_path).name)
    
    def _select_compare_file2(self):
        dpg.add_file_dialog(
            directory_selector=False,
            show=True,
            callback=self._compare_file2_selected,
            file_count=1,
            width=700,
            height=400,
            default_path=str(Path.home() / "Desktop"),
            modal=True
        )
    
    def _compare_file2_selected(self, sender, app_data):
        if app_data and app_data['selections']:
            file_path = list(app_data['selections'].values())[0]
            self.compare_file2_path = file_path
            dpg.set_value("compare_file2_label", Path(file_path).name)
    
    def _select_merge_folder(self):
        dpg.add_file_dialog(
            directory_selector=True,
            show=True,
            callback=self._merge_folder_selected,
            width=700,
            height=400,
            default_path=str(Path.home() / "Desktop"),
            modal=True
        )
    
    def _merge_folder_selected(self, sender, app_data):
        if app_data and app_data['selections']:
            folder_path = list(app_data['selections'].values())[0]
            self.merge_folder_path = folder_path
            dpg.set_value("merge_folder_label", Path(folder_path).name)
    
    def _append_log(self, message):
        current = dpg.get_value("log_text")
        dpg.set_value("log_text", current + message + "\n")
    
    def _update_progress(self, value, message):
        dpg.set_value("progress_bar", value)
        dpg.set_value("progress_label", message)
    
    def _start_clean(self):
        if self.processing:
            return
        
        if not self.clean_file_path:
            self._append_log("[오류] 파일을 선택해주세요.")
            return
        
        self.processing = True
        dpg.configure_item("clean_start_btn", enabled=False)
        
        thread = threading.Thread(target=self._clean_thread, daemon=True)
        thread.start()
    
    def _clean_thread(self):
        try:
            from contact_cleaner.core.processor import ContactProcessor
            from contact_cleaner.utils.file_utils import save_styled_excel
            
            self._update_progress(0.1, "파일 로딩 중...")
            self._append_log(f"[시작] 파일 정리 시작: {Path(self.clean_file_path).name}")
            
            processor = ContactProcessor(self.clean_file_path)
            self._update_progress(0.3, "주소록 처리 중...")
            
            result = processor.process()
            
            self._update_progress(0.7, "파일 저장 중...")
            
            input_path = Path(self.clean_file_path)
            output_path = input_path.parent / f"{input_path.stem}_정리결과.xlsx"
            
            save_styled_excel(result.comparison_data, output_path)
            
            self._update_progress(1.0, "완료")
            self._append_log(f"[완료] 저장 완료")
            self._append_log(f"파일 경로: {output_path}")
            
        except Exception as e:
            self._append_log(f"[오류] {str(e)}")
            self._update_progress(0.0, "오류 발생")
        finally:
            self.processing = False
            dpg.configure_item("clean_start_btn", enabled=True)
    
    def _start_compare(self):
        if self.processing:
            return
        
        if not self.compare_file1_path or not self.compare_file2_path:
            self._append_log("[오류] 두 개의 파일을 모두 선택해주세요.")
            return
        
        self.processing = True
        dpg.configure_item("compare_start_btn", enabled=False)
        
        thread = threading.Thread(target=self._compare_thread, daemon=True)
        thread.start()
    
    def _compare_thread(self):
        try:
            from contact_cleaner.core.processor import ContactProcessor
            from contact_cleaner.utils.file_utils import save_styled_excel
            
            self._update_progress(0.1, "파일 로딩 중...")
            self._append_log(f"[시작] 대조 시작")
            self._append_log(f"파일1: {Path(self.compare_file1_path).name}")
            self._append_log(f"파일2: {Path(self.compare_file2_path).name}")
            
            self._update_progress(0.3, "파일1 처리 중...")
            processor1 = ContactProcessor(self.compare_file1_path)
            result1 = processor1.process()
            
            self._update_progress(0.5, "파일2 처리 중...")
            processor2 = ContactProcessor(self.compare_file2_path)
            result2 = processor2.process()
            
            self._update_progress(0.7, "대조 중...")
            
            phones1 = {item.get('변환됨') for item in result1.transformed_data if item.get('변환됨')}
            
            comparison_data = []
            for item in result2.comparison_data:
                phone = item.get('변환됨') or item.get('원본 전화번호', '')
                if phone in phones1:
                    item['검증'] = 'O'
                else:
                    item['검증'] = 'X'
                comparison_data.append(item)
            
            self._update_progress(0.9, "파일 저장 중...")
            
            file1_name = Path(self.compare_file1_path).stem
            file2_path = Path(self.compare_file2_path)
            output_path = file2_path.parent / f"주소록_{file1_name}_대조결과.xlsx"
            
            save_styled_excel(comparison_data, output_path)
            
            self._update_progress(1.0, "완료")
            self._append_log(f"[완료] 대조 완료")
            self._append_log(f"파일 경로: {output_path}")
            
        except Exception as e:
            self._append_log(f"[오류] {str(e)}")
            self._update_progress(0.0, "오류 발생")
        finally:
            self.processing = False
            dpg.configure_item("compare_start_btn", enabled=True)
    
    def _start_merge(self):
        if self.processing:
            return
        
        if not self.merge_folder_path:
            self._append_log("[오류] 폴더를 선택해주세요.")
            return
        
        self.processing = True
        dpg.configure_item("merge_start_btn", enabled=False)
        
        thread = threading.Thread(target=self._merge_thread, daemon=True)
        thread.start()
    
    def _merge_thread(self):
        try:
            import openpyxl
            
            self._update_progress(0.1, "파일 검색 중...")
            self._append_log(f"[시작] 병합 시작: {Path(self.merge_folder_path).name}")
            
            folder_path = Path(self.merge_folder_path)
            pattern = "주소록_*_대조결과.xlsx"
            files = list(folder_path.glob(pattern))
            
            if not files:
                self._append_log(f"[오류] '{pattern}' 파일을 찾을 수 없습니다.")
                self._update_progress(0.0, "오류 발생")
                return
            
            self._append_log(f"[정보] {len(files)}개 파일 발견")
            
            self._update_progress(0.3, "파일 로딩 중...")
            
            all_data = []
            for file in files:
                wb = openpyxl.load_workbook(file, data_only=True)
                ws = wb.active
                if not ws:
                    wb.close()
                    continue
                headers = [cell.value for cell in ws[2]]
                
                for row in ws.iter_rows(min_row=3, values_only=True):
                    row_dict = {headers[i]: row[i] for i in range(len(headers)) if i < len(row)}
                    all_data.append(row_dict)
                wb.close()
            
            self._update_progress(0.5, "중복 제거 중...")
            
            seen_phones = set()
            unique_data = []
            phone_col = '휴대폰번호'
            
            for item in all_data:
                phone = item.get(phone_col, '')
                if phone and phone not in seen_phones:
                    seen_phones.add(phone)
                    unique_data.append(item)
                elif not phone:
                    unique_data.append(item)
            
            self._update_progress(0.7, "파일 저장 중...")
            
            from contact_cleaner.utils.file_utils import save_styled_excel
            
            comparison_format = []
            for item in unique_data:
                comparison_format.append({
                    '이름': item.get('이름 (휴대폰에 저장될 이름)', ''),
                    '원본 전화번호': '',
                    '변환됨': item.get('휴대폰번호', ''),
                    '검증': item.get('추천3(비교대상과체크)', '')
                })
            
            output_path = folder_path / "병합결과.xlsx"
            save_styled_excel(comparison_format, output_path)
            
            self._update_progress(1.0, "완료")
            self._append_log(f"[완료] 병합 완료 (총 {len(unique_data)}행)")
            self._append_log(f"파일 경로: {output_path}")
            
        except Exception as e:
            self._append_log(f"[오류] {str(e)}")
            self._update_progress(0.0, "오류 발생")
        finally:
            self.processing = False
            dpg.configure_item("merge_start_btn", enabled=True)
    
    def run(self):
        dpg.create_viewport(title="주소록 정리 v1.3.0", width=1150, height=950, resizable=True)
        dpg.setup_dearpygui()
        dpg.show_viewport()
        dpg.set_primary_window("main_window", True)
        dpg.start_dearpygui()
        dpg.destroy_context()


def main():
    app = ContactCleanerDPG()
    app.run()


if __name__ == "__main__":
    main()
