import csv
import os
import shutil
import subprocess
import sys
from pathlib import Path
from typing import NamedTuple
from datetime import datetime


class OutputPaths(NamedTuple):
    folder: Path
    original: Path
    transformed: Path
    comparison: Path


def get_work_folder() -> Path:
    if getattr(sys, "frozen", False):
        base_dir = Path(sys.executable).parent
    else:
        base_dir = Path(__file__).parent.parent.parent

    work_folder = base_dir / "작업결과"
    work_folder.mkdir(exist_ok=True)
    return work_folder


def create_output_structure(filename: str, output_type: str = "", owner_name: str = "") -> Path:
    work_folder = get_work_folder()
    timestamp = datetime.now().strftime("%y%m%d-%H%M")
    stem = Path(filename).stem

    if stem == "정리":
        prefix = "1_"
        output_filename = f"{prefix}{stem}_{owner_name}_{timestamp}.xlsx" if owner_name else f"{prefix}{stem}결과_{timestamp}.xlsx"
    elif stem == "대조":
        output_filename = f"{stem}_{owner_name}_{timestamp}.xlsx" if owner_name else f"{stem}결과_{timestamp}.xlsx"
    else:
        if output_type:
            output_filename = f"{stem}_{output_type}결과_{timestamp}.xlsx"
        else:
            output_filename = f"{stem}결과_{timestamp}.xlsx"
    
    return work_folder / output_filename


def copy_original_file(src: Path, dst: Path) -> None:
    shutil.copy2(src, dst)


def save_transformed_csv(data: list[dict], path: Path) -> None:
    if not data:
        return

    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=["이름", "변환됨"])
        writer.writeheader()
        writer.writerows(data)


def save_comparison_csv(data: list[dict], path: Path) -> None:
    if not data:
        return

    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(
            f, fieldnames=["이름", "원본 전화번호", "변환됨", "검증"]
        )
        writer.writeheader()
        writer.writerows(data)


def save_styled_excel(data: list[dict], path: Path, progress_callback=None) -> None:
    if not data:
        return

    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    ws = wb.active
    if not ws:
        return

    ws.append([])
    ws.cell(row=1, column=1, value="*선택병합 시에는 (해당 열에 \"O\" 표시해 주세요.)")

    headers = [
        "연번",
        "이름 (휴대폰에 저장될 이름)",
        "휴대폰번호",
        "추천1(DW)",
        "추천2",
        "추천3(비교대상과체크)",
        "선택병합",
        "휴대폰 저장파일명",
    ]
    ws.append(headers)

    gray_fill = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid")
    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    header_font = Font(name="Malgun Gothic", size=11, bold=False)
    header_font_red = Font(name="Malgun Gothic", size=11, bold=False, color="FF0000")
    data_font = Font(name="Malgun Gothic", size=10)
    center_alignment = Alignment(horizontal="center", vertical="center")
    thin_border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    for col_idx in range(1, 9):
        cell = ws.cell(row=2, column=col_idx)
        if col_idx == 7:
            cell.font = header_font_red
        else:
            cell.font = header_font
        cell.alignment = center_alignment
        cell.border = thin_border
        cell.fill = gray_fill

    total_rows = len(data)
    current_row = 3
    for idx, item in enumerate(data, start=1):
        if progress_callback and idx % 500 == 0:
            progress_callback(idx, total_rows)
        
        name_value = item.get("이름", "") or " "
        phone_value = item.get("변환됨") or item.get("원본 전화번호") or " "
        verification = item.get("검증", "")
        
        c1 = ws.cell(row=current_row, column=1, value=idx)
        c2 = ws.cell(row=current_row, column=2, value=name_value)
        c3 = ws.cell(row=current_row, column=3, value=phone_value)
        c4 = ws.cell(row=current_row, column=4)
        c5 = ws.cell(row=current_row, column=5)
        c6 = ws.cell(row=current_row, column=6, value=verification)
        c7 = ws.cell(row=current_row, column=7)
        c8 = ws.cell(row=current_row, column=8)

        c1.border = c2.border = c3.border = c4.border = thin_border
        c5.border = c6.border = c7.border = c8.border = thin_border
        c1.font = c2.font = c3.font = c4.font = data_font
        c5.font = c6.font = c7.font = c8.font = data_font
        c1.alignment = center_alignment
        c6.alignment = center_alignment
        c6.fill = yellow_fill
        
        current_row += 1
    
    if progress_callback:
        progress_callback(total_rows, total_rows)

    widths = {1: 8, 2: 30, 3: 20, 4: 10, 5: 10, 6: 25, 7: 10, 8: 25}
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    if path.suffix != ".xlsx":
        path = path.with_suffix(".xlsx")
    wb.save(path)


def open_folder_in_explorer(path: Path) -> None:
    folder = path if path.is_dir() else path.parent

    if sys.platform == "win32":
        os.startfile(folder)
    elif sys.platform == "darwin":
        subprocess.run(["open", str(folder)])
    else:
        subprocess.run(["xdg-open", str(folder)])
